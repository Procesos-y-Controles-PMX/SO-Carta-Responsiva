#!/usr/bin/env python3
import argparse
import json
import os
import unicodedata
import urllib.error
import urllib.parse
import urllib.request

from openpyxl import load_workbook


def clean(value):
    return " ".join(str(value or "").strip().split())


def normalized(value):
    decomposed = unicodedata.normalize("NFD", clean(value))
    return "".join(char for char in decomposed if not unicodedata.combining(char)).upper()


def phone(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean(value)


def rows_from_sheet(path, sheet_name=None):
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.active
    return list(sheet.iter_rows(values_only=True))


def api_request(path, method="GET", payload=None, prefer=None):
    base_url = os.environ["NEXT_PUBLIC_SUPABASE_URL"].rstrip("/")
    service_key = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    headers = {
        "apikey": service_key,
        "Authorization": f"Bearer {service_key}",
        "Content-Type": "application/json",
    }
    if prefer:
        headers["Prefer"] = prefer
    body = json.dumps(payload).encode() if payload is not None else None
    request = urllib.request.Request(
        f"{base_url}/rest/v1/{path}",
        data=body,
        headers=headers,
        method=method,
    )
    try:
        with urllib.request.urlopen(request) as response:
            content = response.read()
            return json.loads(content) if content else None
    except urllib.error.HTTPError as error:
        detail = error.read().decode()
        raise RuntimeError(f"{method} {path}: {error.code} {detail}") from error


def api_get_all(path: str):
    """Paginate through PostgREST results (default page size 1000)."""
    rows = []
    offset = 0
    page = 1000
    while True:
        sep = "&" if "?" in path else "?"
        chunk = api_request(f"{path}{sep}limit={page}&offset={offset}")
        if not chunk:
            break
        rows.extend(chunk)
        if len(chunk) < page:
            break
        offset += page
    return rows


def dedupe_managers(managers):
    """Collapse repeats in file by (centro, normalized nombre); last wins."""
    by_key = {}
    collapsed = 0
    for manager in managers:
        key = f'{manager["codigo_sap"]}:{normalized(manager["nombre"])}'
        if key in by_key:
            collapsed += 1
        nombre = normalized(manager["nombre"])
        by_key[key] = {
            "codigo_sap": manager["codigo_sap"],
            "nombre": nombre,
            "nombre_normalizado": nombre,
        }
    return list(by_key.values()), collapsed


def parse_branches(path):
    rows = rows_from_sheet(path)
    branches = []
    managers = []
    for row in rows[1:]:
        branch_name = clean(row[0] if len(row) > 0 else "")
        center = clean(row[1] if len(row) > 1 else "").upper()
        if not branch_name or not center:
            continue
        manager_name = clean(row[3] if len(row) > 3 else "")
        manager_email = clean(row[5] if len(row) > 5 else "").lower()
        branches.append(
            {
                "nombre": branch_name,
                "codigo_sap": center,
                "prefijo_folio": center,
                "region": clean(row[2] if len(row) > 2 else "") or None,
                "gerente_nombre": manager_name or None,
                "gerente_celular": phone(row[4] if len(row) > 4 else "") or None,
                "gerente_email": manager_email or None,
                "direccion": clean(row[6] if len(row) > 6 else "") or None,
                "estado": clean(row[7] if len(row) > 7 else "") or None,
                "activo": normalized(row[8] if len(row) > 8 else "Activo") == "ACTIVO",
            }
        )
        if manager_name:
            managers.append({"codigo_sap": center, "nombre": manager_name})
    return branches, managers


ROLE_MAP = {
    "USUARIO": "usuario",
    "ADMINSITRADOR ZONA": "administrador_zona",
    "ADMINISTRADOR ZONA": "administrador_zona",
    "ADMINISTRADOR GENERAL": "administrador_general",
}


def parse_access(path):
    rows = rows_from_sheet(path, "Accesos")
    users = []
    missing_email = []
    unknown_roles = []
    seen_emails = set()

    for row in rows[1:]:
        name = clean(row[0] if len(row) > 0 else "")
        email = clean(row[1] if len(row) > 1 else "").lower()
        center = clean(row[2] if len(row) > 2 else "").upper()
        region = clean(row[4] if len(row) > 4 else "") or None
        raw_role = normalized(row[5] if len(row) > 5 else "")
        role = ROLE_MAP.get(raw_role)
        if not email:
            missing_email.append(name or "(sin nombre)")
            continue
        if not role:
            unknown_roles.append({"email": email, "rol": raw_role})
            continue
        if email in seen_emails:
            continue
        seen_emails.add(email)
        users.append(
            {
                "nombre_completo": name or None,
                "email": email,
                "codigo_sap": center or None,
                "region": region,
                "rol": role,
            }
        )
    return users, missing_email, unknown_roles


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("branches_workbook")
    parser.add_argument("access_workbook")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    branches, managers = parse_branches(args.branches_workbook)
    managers, managers_collapsed = dedupe_managers(managers)
    users, missing_email, unknown_roles = parse_access(args.access_workbook)
    role_counts = {
        role: sum(user["rol"] == role for user in users)
        for role in ROLE_MAP.values()
    }

    print(f"Sucursales preparadas: {len(branches)}")
    print(f"Gerentes preparados como responsables: {len(managers)}")
    if managers_collapsed:
        print(f"Gerentes duplicados colapsados en archivo: {managers_collapsed}")
    print(f"Accesos válidos: {len(users)} {role_counts}")
    print(f"Filas sin correo: {len(missing_email)}")
    print(f"Roles desconocidos: {len(unknown_roles)}")
    if unknown_roles:
        print(json.dumps(unknown_roles, ensure_ascii=False))

    if args.dry_run:
        return
    if not os.environ.get("NEXT_PUBLIC_SUPABASE_URL") or not os.environ.get(
        "SUPABASE_SERVICE_ROLE_KEY"
    ):
        raise RuntimeError(
            "Faltan NEXT_PUBLIC_SUPABASE_URL o SUPABASE_SERVICE_ROLE_KEY."
        )

    existing_branches = api_get_all("cr_sucursales?select=id,codigo_sap")
    existing_branch_by_center = {
        clean(branch["codigo_sap"]).upper(): branch
        for branch in existing_branches
        if clean(branch.get("codigo_sap"))
    }
    for branch in branches:
        existing_branch = existing_branch_by_center.get(branch["codigo_sap"])
        if existing_branch:
            api_request(
                f'cr_sucursales?id=eq.{existing_branch["id"]}',
                method="PATCH",
                payload=branch,
                prefer="return=minimal",
            )
        else:
            api_request(
                "cr_sucursales",
                method="POST",
                payload=branch,
                prefer="return=minimal",
            )
    stored_branches = api_get_all(
        "cr_sucursales?select=id,codigo_sap,region&activo=eq.true"
    )
    branch_by_center = {
        clean(branch["codigo_sap"]).upper(): branch for branch in stored_branches
    }

    existing_responsibles = api_get_all(
        "cr_responsables?select=id,id_sucursal,nombre,nombre_normalizado"
    )
    existing_by_key = {}
    for item in existing_responsibles:
        key_name = item.get("nombre_normalizado") or normalized(item["nombre"])
        existing_by_key[f'{item["id_sucursal"]}:{key_name}'] = item

    manager_inserts = []
    manager_updates = 0
    for manager in managers:
        branch = branch_by_center.get(manager["codigo_sap"])
        if not branch:
            continue
        key = f'{branch["id"]}:{manager["nombre_normalizado"]}'
        existing = existing_by_key.get(key)
        if existing:
            api_request(
                f'cr_responsables?id=eq.{existing["id"]}',
                method="PATCH",
                payload={
                    "nombre": manager["nombre"],
                    "nombre_normalizado": manager["nombre_normalizado"],
                    "activo": True,
                },
                prefer="return=minimal",
            )
            manager_updates += 1
        else:
            manager_inserts.append(
                {
                    "id_sucursal": branch["id"],
                    "nombre": manager["nombre"],
                    "nombre_normalizado": manager["nombre_normalizado"],
                    "activo": True,
                }
            )
            existing_by_key[key] = {"id": "(pending)", **manager_inserts[-1]}
    if manager_inserts:
        for index in range(0, len(manager_inserts), 200):
            api_request(
                "cr_responsables",
                method="POST",
                payload=manager_inserts[index : index + 200],
                prefer="return=minimal",
            )

    user_payload = []
    missing_centers = []
    for user in users:
        branch = (
            branch_by_center.get(user["codigo_sap"]) if user["codigo_sap"] else None
        )
        if user["rol"] == "usuario" and not branch:
            missing_centers.append(
                {"email": user["email"], "centro": user["codigo_sap"]}
            )
            continue
        user_payload.append(
            {
                "email": user["email"],
                "nombre_completo": user["nombre_completo"],
                "rol": user["rol"],
                "id_sucursal": branch["id"] if branch else None,
                "region": user["region"],
                "activo": True,
            }
        )

    if missing_centers:
        print("Accesos omitidos por centro inexistente:")
        print(json.dumps(missing_centers, ensure_ascii=False))

    api_request(
        "cr_usuarios?activo=eq.true",
        method="PATCH",
        payload={"activo": False},
        prefer="return=minimal",
    )
    api_request(
        "cr_usuarios?on_conflict=email",
        method="POST",
        payload=user_payload,
        prefer="resolution=merge-duplicates,return=minimal",
    )

    shared_accounts = api_request(
        "ctz_usuarios?select=email&activo=eq.true&password=not.is.null"
    )
    shared_emails = {
        clean(account["email"]).lower()
        for account in shared_accounts
        if clean(account.get("email"))
    }
    missing_shared_login = [
        {"nombre": user["nombre_completo"], "email": user["email"]}
        for user in user_payload
        if user["email"] not in shared_emails
    ]

    print(f"Sucursales sincronizadas: {len(branches)}")
    print(
        "Gerentes responsables: "
        f"{manager_updates} existentes, {len(manager_inserts)} nuevos"
    )
    print(f"Accesos sincronizados: {len(user_payload)}")
    print(f"Accesos sin cuenta compartida: {len(missing_shared_login)}")
    if missing_shared_login:
        print(json.dumps(missing_shared_login, ensure_ascii=False))


if __name__ == "__main__":
    main()
